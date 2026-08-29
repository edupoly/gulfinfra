from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "deliverables" / "GulfInfraHub_Client_User_Guide.docx"

# compact_reference_guide preset with GulfInfraHub brand overrides.
NAVY, BLUE, AMBER = "0B1F3A", "1D4ED8", "F4B400"
INK, MUTED, LIGHT = "162235", "64748B", "E8EEF5"
PALE_BLUE, PALE_AMBER, WHITE = "EAF2FF", "FFF7D6", "FFFFFF"


def set_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size, run.bold, run.italic = Pt(size), bold, italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def cell_width(cell, width):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:tcW"))
    if node is None:
        node = OxmlElement("w:tcW")
        props.append(node)
    node.set(qn("w:w"), str(width))
    node.set(qn("w:type"), "dxa")


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    table.autofit = False
    props = table._tbl.tblPr
    total = props.find(qn("w:tblW"))
    if total is None:
        total = OxmlElement("w:tblW")
        props.append(total)
    total.set(qn("w:w"), str(sum(widths)))
    total.set(qn("w:type"), "dxa")
    indent = props.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        props.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell_width(cell, widths[index])
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def page_number(paragraph):
    run = paragraph.add_run()
    begin, separate, end = (OxmlElement("w:fldChar") for _ in range(3))
    begin.set(qn("w:fldCharType"), "begin")
    separate.set(qn("w:fldCharType"), "separate")
    end.set(qn("w:fldCharType"), "end")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    value = OxmlElement("w:t")
    value.text = "1"
    for node in (begin, instruction, separate, value, end):
        run._r.append(node)
    set_font(run, 8.5, color=MUTED)


def para(doc, text, size=11, bold=False, color=INK, after=6, italic=False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size, bold, color, italic)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text))


def step(doc, title, detail):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(f"{title} — "), bold=True, color=NAVY)
    set_font(p.add_run(detail))


def callout(doc, label, text, tone="blue"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, PALE_BLUE if tone == "blue" else PALE_AMBER)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(f"{label}: "), 10.5, True, NAVY)
    set_font(p.add_run(text), 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def label_table(doc, rows, widths=(2700, 6660), header=None):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    if header:
        cells = table.add_row().cells
        for i, value in enumerate(header):
            shade(cells[i], NAVY)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            set_font(cells[i].paragraphs[0].add_run(value), 10, True, WHITE)
    for left, right in rows:
        cells = table.add_row().cells
        shade(cells[0], LIGHT)
        for i, value in enumerate((left, right)):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            set_font(p.add_run(value), 10, i == 0, NAVY if i == 0 else INK)
    table_geometry(table, list(widths))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def module_intro(doc, purpose, audience, entry):
    label_table(doc, [("Purpose", purpose), ("Primary users", audience), ("Main entry point", entry)])


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(11)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before, normal.paragraph_format.space_after = Pt(0), Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7), ("Heading 3", 12, NAVY, 10, 5)):
    style = doc.styles[name]
    style.font.name, style.font.size, style.font.bold = "Calibri", Pt(size), True
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
    style.paragraph_format.keep_with_next = True
for name in ("List Bullet", "List Bullet 2", "List Number"):
    style = doc.styles[name]
    style.font.name, style.font.size = "Calibri", Pt(11)
    style.paragraph_format.space_after, style.paragraph_format.line_spacing = Pt(4), 1.25

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("GulfInfraHub  |  Client User Guide"), 8.5, True, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer.add_run("Client reference  •  August 2026  |  "), 8.5, color=MUTED)
page_number(footer)

# Editorial cover.
para(doc, "CLIENT USER GUIDE", 10, True, AMBER, 22, align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "GulfInfraHub", 30, True, NAVY, 5, align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "Authentication, authorization and marketplace modules", 15, color=BLUE, after=24, align=WD_ALIGN_PARAGRAPH.CENTER)
callout(doc, "Purpose", "A clear operational reference for client users, marketplace participants and administrators using the current application.")
heading(doc, "How this guide is organized")
label_table(doc, [("Section 1", "Project orientation and common navigation"), ("Section 2", "Authentication and authorization in one consolidated section"), ("Sections 3–8", "One independent section for each core marketplace module"), ("Sections 9–11", "Account workspace, administration, support and acceptance checks")], header=("Guide area", "Coverage"))
heading(doc, "Core modules")
para(doc, "Contractors & Industrial Services  •  Projects & Tenders  •  RFQ Marketplace  •  Equipment Marketplace  •  Construction & Industrial Materials  •  Business Opportunities", 11, True, NAVY, 8)
callout(doc, "Scope", "This document describes functionality present in the supplied project. Items such as memberships, payments, banners and blog management are not represented as active modules in this guide.", "amber")

page_break(doc)
heading(doc, "Contents")
label_table(doc, [("1", "Project overview and navigation"), ("2", "Authentication and authorization"), ("3", "Module: Contractors & Industrial Services"), ("4", "Module: Projects & Tenders"), ("5", "Module: RFQ Marketplace"), ("6", "Module: Equipment Marketplace"), ("7", "Module: Construction & Industrial Materials"), ("8", "Module: Business Opportunities"), ("9", "Account workspace and shared features"), ("10", "Administration module"), ("11", "Support, privacy and client acceptance")], widths=(900, 8460), header=("Section", "Subject"))
heading(doc, "Reading pattern for each module", 2)
for item in ("Purpose and intended users", "What users can view and do", "Listing or transaction workflow", "Statuses, access rules and expected results"):
    bullet(doc, item)

page_break(doc)
heading(doc, "1. Project overview and navigation")
para(doc, "GulfInfraHub is a GCC-focused construction and industrial marketplace. It connects buyers, suppliers, contractors, project owners and opportunity publishers through searchable listings and controlled transaction workflows.")
heading(doc, "Primary navigation", 2)
label_table(doc, [("Home", "Featured and latest marketplace content across the platform"), ("Marketplace directories", "Contractors, projects, RFQs, equipment, materials and business opportunities"), ("Search", "Unified search across all six core modules"), ("Add listing", "Guided creation flow for the five listing-based modules"), ("Account", "Owned listings, saved items, RFQs, quotations, applications, messages and profile"), ("Admin", "Restricted management area for authorized administrators")], header=("Area", "Purpose"))
heading(doc, "Common listing lifecycle", 2)
label_table(doc, [("Draft", "Work is saved but not submitted for public review"), ("Pending", "Submitted and awaiting administrator moderation"), ("Published", "Visible in the public marketplace"), ("On hold", "Temporarily paused by an administrator"), ("Rejected", "Declined during moderation"), ("Closed", "No longer active for responses")], header=("Status", "Meaning"))
callout(doc, "Visibility rule", "Public directory and detail pages are intended to show published records. Owners and administrators can access additional records through their workspaces.")

page_break(doc)
heading(doc, "2. Authentication and authorization")
para(doc, "Authentication confirms who the user is. Authorization determines what that signed-in user is allowed to view or change. GulfInfraHub combines email verification, password login, database-backed sessions, account status checks, ownership rules and administrator role checks.")
heading(doc, "2.1 Authentication methods", 2)
label_table(doc, [("Email OTP", "A four-digit code is emailed to the user. It expires after 10 minutes, is single-use, and allows up to five incorrect attempts."), ("Password", "Created after first successful email verification; minimum eight characters with at least one letter and one number."), ("Password reset", "A single-use email reset link is available from Forgot password and is valid for 30 minutes."), ("Session", "A random session token is stored in an HTTP-only cookie; only its hash is stored in the database. The session lasts up to 30 days.")], header=("Control", "Implementation and user effect"))
heading(doc, "2.2 Sign-in and first-time activation", 2)
for title, detail in (("Enter email", "Open Sign in and provide the assigned email address."), ("Request the code", "Select Email me a code. Requests are rate-limited to one per minute."), ("Verify", "Enter the newest four-digit code within 10 minutes."), ("Create a password", "First-time users create and confirm a compliant password."), ("Continue", "Administrators are directed to Admin; normal users are directed to My listings.")):
    step(doc, title, detail)
heading(doc, "2.3 Authorization model", 2)
label_table(doc, [("Public visitor", "Browse published directories and details, search, view public pages and submit the Contact form."), ("Signed-in user", "Use account features and create marketplace records. Access is limited to records owned by that account where ownership applies."), ("Blocked user", "May be identifiable by the system but is prevented from protected marketplace and account activity."), ("Administrator", "Access admin pages, moderate submitted listings, feature published records, review users/applications/contact messages, and block or unblock normal users.")], header=("Role/state", "Authorized access"))
callout(doc, "Security boundary", "The interface is not the only control. Protected server actions verify the current session, user role, blocked state and/or record ownership before changing data.", "amber")

page_break(doc)
heading(doc, "2.4 Authorization by action", 2)
label_table(doc, [("Create a listing", "Signed-in, active user; listing publication also uses email verification in the creation flow."), ("Edit/delete a listing", "The listing owner; administrator moderation remains separate."), ("Save a listing", "Signed-in user."), ("Create an RFQ", "Active user with a fresh RFQ-purpose OTP."), ("Submit a quotation", "Signed-in supplier account; the buyer cannot quote on its own RFQ."), ("Apply to a project", "Signed-in applicant; project ownership and application ownership are checked for later actions."), ("Moderate/feature", "Administrator only; only published listings can be featured."), ("Block/unblock users", "Administrator only; administrators cannot block themselves or another administrator through this interface.")], header=("Action", "Required authorization"))
heading(doc, "2.5 Account security operations", 2)
for item in ("Changing a password requires the current password when one already exists.", "After a password change, other active sessions are removed.", "Signing out removes the current database session and cookie.", "Passwords are stored as salted scrypt hashes; OTPs and session tokens are stored as hashes."):
    bullet(doc, item)
heading(doc, "Common access issues", 2)
label_table(doc, [("OTP not received", "Check Spam/Junk and the recipient address; wait one minute before another request."), ("OTP rejected", "Use the newest code and request another if it is older than 10 minutes or has reached five failed attempts."), ("Authorization error", "Confirm the correct account is signed in and that it owns the record or has the required admin role."), ("Blocked message", "Contact an administrator; protected activity remains unavailable until the account is restored.")], header=("Issue", "Recommended action"))

page_break(doc)
heading(doc, "3. Module: Contractors & Industrial Services")
module_intro(doc, "Publish and discover contractor and industrial-service company profiles.", "Public visitors, contractors, buyers and administrators", "/contractors and Add listing → Contractor")
heading(doc, "Directory and detail experience", 2)
for item in ("Browse and search contractor profiles; filter by supported location information.", "View company identity, type, description, services, countries/cities served and contact channels.", "Review logo, gallery, licences, documents and featured projects when supplied.", "Save a listing from supported listing cards/details while signed in."):
    bullet(doc, item)
heading(doc, "Create and manage a contractor profile", 2)
for title, detail in (("Select Contractor", "Open Add listing and choose the contractor listing type."), ("Enter company information", "Provide name, company type, service categories, establishment details and description."), ("Add coverage and contacts", "Provide countries, cities, areas served, website, email, phone, WhatsApp and address as applicable."), ("Add proof and media", "Upload authorized logo, gallery images, licences, documents and project references."), ("Review and submit", "Verify the summary and complete the publication authorization step.")):
    step(doc, title, detail)
heading(doc, "Control and result", 2)
para(doc, "The creator becomes the listing owner. The owner manages the record in My listings; an administrator approves, rejects or places a submitted record on hold. A published contractor may be marked as featured by an administrator.")

page_break(doc)
heading(doc, "4. Module: Projects & Tenders")
module_intro(doc, "Publish project/tender opportunities and manage applications from interested companies.", "Project owners, applicants, public visitors and administrators", "/projects-tenders and Add listing → Project / Tender")
heading(doc, "Listing capabilities", 2)
for item in ("Search and browse projects and tenders by relevant marketplace data.", "Show title, type, status, summary, description, budget/value, deadline, client, location, sectors and supporting documents.", "Open an application flow from an eligible published project/tender detail page."):
    bullet(doc, item)
heading(doc, "Applicant workflow", 2)
for title, detail in (("Open a project", "Review the description, deadline, requirements and files."), ("Select Apply", "Sign in if required and complete company, contact and proposed-role information."), ("Submit", "The application is associated with the applicant account and project."), ("Track", "Use My applications to view status and timeline; edit while submitted or withdraw while still active.")):
    step(doc, title, detail)
heading(doc, "Project-owner workflow", 2)
for item in ("Use Received applications to review applicants for owned project listings.", "Open application details and move eligible applications to shortlisted, accepted or rejected.", "Application statuses include submitted, under review, shortlisted, accepted, rejected and withdrawn."):
    bullet(doc, item)
callout(doc, "Ownership", "Applicants manage their own applications; project owners review applications received for projects they own; administrators can inspect applications across the platform.")

page_break(doc)
heading(doc, "5. Module: RFQ Marketplace")
module_intro(doc, "Allow buyers to request quotations and suppliers to submit, discuss and track commercial offers.", "Buyers, suppliers and administrators", "/rfqs, My RFQs and My Quotations")
heading(doc, "Buyer workflow", 2)
for title, detail in (("Start an RFQ", "Provide contact details and request a fresh RFQ-purpose email code."), ("Describe the requirement", "Enter project, category, material/service, location, dates, quantity, unit, specifications, address, budget and notes."), ("Attach documents", "Add BOQ, drawings, specifications or other relevant files."), ("Review and publish", "Verify the details and consume the one-time authorization code."), ("Manage responses", "Use My RFQs to search owned RFQs, review quotations and control the RFQ lifecycle.")):
    step(doc, title, detail)
heading(doc, "Supplier workflow", 2)
for title, detail in (("Browse RFQs", "Open a relevant request and review its requirements and closing date."), ("Submit a quotation", "Provide supplier/company details, pricing, delivery time, validity, warranty, payment terms, remarks and quotation file where applicable."), ("Track the offer", "Use My Quotations to view, edit eligible offers or withdraw an active submission."), ("Continue the conversation", "Use the quotation workspace for status history, messages and files.")):
    step(doc, title, detail)
heading(doc, "Quotation outcomes", 2)
label_table(doc, [("Draft / Submitted", "Supplier is preparing or has submitted the quotation."), ("Under review / Shortlisted", "Buyer is evaluating or has shortlisted the offer."), ("Awarded", "Supplier can continue the awarded workflow and download the award letter where available."), ("Rejected / Withdrawn", "Buyer declined the quotation or supplier withdrew it.")], header=("Status", "Meaning"))

page_break(doc)
heading(doc, "6. Module: Equipment Marketplace")
module_intro(doc, "Advertise and discover construction or industrial equipment offered for sale, rent or wanted.", "Equipment owners, rental/sales companies, buyers and administrators", "/equipment-marketplace and Add listing → Equipment")
heading(doc, "Directory and listing information", 2)
for item in ("Browse and search equipment records, including featured items.", "Identify listing intent such as sale, rent or wanted through the listing data.", "Review equipment type, title, description, brand, model, year, condition, price, location and availability information.", "View images and seller contact channels supplied with the record."):
    bullet(doc, item)
heading(doc, "Publishing workflow", 2)
for title, detail in (("Choose Equipment", "Open the listing-type selector."), ("Describe the asset", "Provide category/type, commercial terms, specifications, condition and location."), ("Add media and contacts", "Upload authorized images and provide accurate contact information."), ("Review and submit", "Complete the authorization step and send the listing for moderation.")):
    step(doc, title, detail)
heading(doc, "Control and result", 2)
para(doc, "The owner can access the equipment record from My listings. Administrators control publication status and may feature a published equipment listing. Rejected or on-hold equipment cannot remain featured.")

page_break(doc)
heading(doc, "7. Module: Construction & Industrial Materials")
module_intro(doc, "Connect suppliers and buyers around construction and industrial material listings.", "Material suppliers, procurement teams, buyers and administrators", "/construction-materials and Add listing → Material")
heading(doc, "Directory and listing information", 2)
for item in ("Browse/search material records and open individual material details.", "Review material name, category/type, description, specifications, unit, price/terms, availability, country and city.", "View supplier identity, contact methods, images and supporting documents where provided."):
    bullet(doc, item)
heading(doc, "Publishing workflow", 2)
for title, detail in (("Choose Material", "Open Add listing and select the materials flow."), ("Enter the offering", "Add the material identity, category, description, specification and commercial details."), ("Add location and supplier contacts", "Provide accurate marketplace and communication details."), ("Upload media/documents", "Use files the publisher is authorized to distribute."), ("Review and submit", "Verify the final summary and complete publication authorization.")):
    step(doc, title, detail)
heading(doc, "Related procurement path", 2)
para(doc, "A buyer needing competitive supplier responses can use the RFQ module instead of relying only on direct listing contact. This keeps a formal request and received quotations in the buyer workspace.")

page_break(doc)
heading(doc, "8. Module: Business Opportunities")
module_intro(doc, "Publish and discover businesses for sale, businesses wanted and investment opportunities.", "Business owners, investors, opportunity seekers and administrators", "/business-opportunities and Add listing → Business Opportunity")
heading(doc, "Directory and listing information", 2)
for item in ("Browse and search opportunity records and view individual details.", "Review opportunity type, category, title, business/summary information, asking price or investment data, location and contact details.", "View images/documents provided for the opportunity and save relevant records while signed in."):
    bullet(doc, item)
heading(doc, "Publishing workflow", 2)
for title, detail in (("Choose Business Opportunity", "Open the corresponding listing flow."), ("Describe the opportunity", "State the opportunity type, business details, value/asking price and location clearly."), ("Add contact and media", "Provide accurate communication details and authorized supporting content."), ("Review and submit", "Complete the email authorization and send the record for moderation.")):
    step(doc, title, detail)
callout(doc, "Content quality", "Do not publish confidential financial information, personal documents or claims that cannot be supported. Administrators should review opportunity descriptions and attachments before approval.", "amber")

page_break(doc)
heading(doc, "9. Account workspace and shared features")
para(doc, "The account workspace brings together records and activity associated with the signed-in user.")
label_table(doc, [("My listings", "View owned contractor, project, equipment, material and business-opportunity records; open eligible records for editing or deletion."), ("Saved listings", "Review marketplace items saved by the account."), ("My RFQs", "Buyer workspace for requests created by the account and quotations received."), ("My Quotations", "Supplier workspace for submitted offers, status and follow-up."), ("My Applications", "Applicant view of submitted project/tender applications and timelines."), ("Received Applications", "Project-owner view of applications received for owned projects."), ("Messages", "Conversations connected to RFQ quotations."), ("Notifications", "RFQ/quotation status activity relevant to the user."), ("Profile & security", "Personal/company details, profile image and password management.")], header=("Workspace", "Purpose"))
heading(doc, "Global search", 2)
para(doc, "Unified Global Search covers contractors, projects, RFQs, equipment, materials and business opportunities. Users can combine a keyword with module/category and supported location/type filters, then open the matching public record.")
heading(doc, "Contact form", 2)
para(doc, "Visitors can send a general enquiry from Contact Us. Submissions are stored for administrator review. Contact information and message content should be treated as confidential business data.")

page_break(doc)
heading(doc, "10. Administration module")
module_intro(doc, "Provide platform-level visibility and controlled moderation for authorized administrators.", "Users whose account role is admin", "/admin")
heading(doc, "Administration areas", 2)
label_table(doc, [("Overview", "Marketplace totals, registered users, blocked users and listing-status summaries."), ("Categories", "Counts by marketplace category and listing status."), ("Users", "Review non-admin accounts, verification state and listing counts; block or unblock access."), ("All listings", "Review records across all six modules, moderate submitted records and feature eligible published listings."), ("Project applications", "Filter and inspect applications submitted across projects and tenders."), ("Contact submissions", "Read public contact messages and use supplied contact channels."), ("Security", "Change the administrator password and invalidate other sessions.")], header=("Admin area", "Responsibility"))
heading(doc, "Moderation workflow", 2)
for title, detail in (("Locate the record", "Confirm listing type, title/reference, owner and current status."), ("Inspect content", "Check business details, location, contacts, images and documents."), ("Decide", "Approve to publish, reject, or place the record on hold."), ("Feature if appropriate", "Only a published non-RFQ listing can be featured."), ("Verify", "Confirm the admin status badge, category totals and public visibility.")):
    step(doc, title, detail)
callout(doc, "Admin safeguard", "Every administrator mutation performs a server-side admin-role check. User blocking also prevents targeting the current administrator or another admin account.")

page_break(doc)
heading(doc, "11. Support, privacy and client acceptance")
heading(doc, "Safe operating practices", 2)
for item in ("Never share passwords, OTPs, reset links or session cookies.", "Use clearly labelled test records and remove or close them after acceptance testing.", "Do not upload real confidential drawings, financial statements or identity documents during testing.", "Confirm the recipient and business need before downloading or forwarding application and quotation files.", "Sign out after using a shared device."):
    bullet(doc, item)
heading(doc, "Client acceptance checklist", 2)
label_table(doc, [("☐", "Complete OTP activation, create a password, sign out and sign in again."), ("☐", "Confirm a normal user cannot open administrator pages or moderate records."), ("☐", "Create one clearly marked test listing in each required marketplace module."), ("☐", "Approve one listing and confirm its public visibility and search result."), ("☐", "Create an RFQ, submit a supplier quotation and exercise the status/message workflow."), ("☐", "Submit a project application and review it from the project-owner workspace."), ("☐", "Block a normal test user, verify restricted activity, then restore the account."), ("☐", "Submit a Contact Us message and verify it appears in Admin."), ("☐", "Change the admin password and confirm other sessions are signed out.")], widths=(600, 8760), header=("Done", "Acceptance test"))
heading(doc, "Issue reporting", 2)
para(doc, "Include the page or module, signed-in role, action performed, expected result, actual result, date/time and a screenshot. Do not include passwords, OTPs, reset links or confidential uploaded files.")
callout(doc, "Document basis", "Prepared from the current GulfInfraHub application structure and implemented workflows as of August 2026.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
