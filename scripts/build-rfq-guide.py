from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "deliverables" / "GulfInfraHub_RFQ_Feature_Guide.docx"
NAVY, BLUE, AMBER, INK, MUTED = "0B1F3A", "1D4ED8", "F4B400", "162235", "64748B"
LIGHT, PALE_BLUE, PALE_AMBER, WHITE = "E8EEF5", "EAF2FF", "FFF7D6", "FFFFFF"

def font(run, size=11, bold=False, color=INK):
    run.font.name = "Calibri"; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri"); run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size, run.bold, run.font.color.rgb = Pt(size), bold, RGBColor.from_string(color)

def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr(); node = props.find(qn("w:shd")) or OxmlElement("w:shd")
    if node.getparent() is None: props.append(node)
    node.set(qn("w:fill"), fill)

def table_geometry(table, widths):
    table.autofit = False; props = table._tbl.tblPr
    for tag, value in (("tblW", sum(widths)), ("tblInd", 120)):
        node = props.find(qn(f"w:{tag}")) or OxmlElement(f"w:{tag}")
        if node.getparent() is None: props.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            props = cell._tc.get_or_add_tcPr(); tcw = props.find(qn("w:tcW")) or OxmlElement("w:tcW")
            if tcw.getparent() is None: props.append(tcw)
            tcw.set(qn("w:w"), str(widths[i])); tcw.set(qn("w:type"), "dxa")
            margins = props.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
            if margins.getparent() is None: props.append(margins)
            for side, value in (("top",80),("start",120),("bottom",80),("end",120)):
                node = margins.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}")
                if node.getparent() is None: margins.append(node)
                node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def para(doc, text, size=11, bold=False, color=INK, after=6, align=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.25
    if align is not None: p.alignment=align
    font(p.add_run(text), size, bold, color)

def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.keep_with_next=True; p.add_run(text); return p

def bullet(doc, title, detail=""):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.25
    font(p.add_run(title + (" — " if detail else "")), 11, True, NAVY)
    if detail: font(p.add_run(detail))

def callout(doc, label, text, amber=False):
    t=doc.add_table(rows=1, cols=1); t.style="Table Grid"; table_geometry(t,[9360]); c=t.cell(0,0); shade(c, PALE_AMBER if amber else PALE_BLUE)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.2; font(p.add_run(label+": "),10.5,True,NAVY); font(p.add_run(text),10.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def grid(doc, rows, widths=(2600,6760), header=("Item","Behaviour")):
    t=doc.add_table(rows=0, cols=2); t.style="Table Grid"; cells=t.add_row().cells
    for i,text in enumerate(header): shade(cells[i],NAVY); font(cells[i].paragraphs[0].add_run(text),10,True,WHITE)
    for left,right in rows:
        cells=t.add_row().cells; shade(cells[0],LIGHT); font(cells[0].paragraphs[0].add_run(left),10,True,NAVY); font(cells[1].paragraphs[0].add_run(right),10)
        for c in cells: c.paragraphs[0].paragraph_format.space_after=Pt(0)
    table_geometry(t,list(widths)); doc.add_paragraph().paragraph_format.space_after=Pt(1)

def page_number(p):
    run=p.add_run(); begin,sep,end=(OxmlElement("w:fldChar") for _ in range(3)); begin.set(qn("w:fldCharType"),"begin"); sep.set(qn("w:fldCharType"),"separate"); end.set(qn("w:fldCharType"),"end")
    inst=OxmlElement("w:instrText"); inst.set(qn("xml:space"),"preserve"); inst.text=" PAGE "; value=OxmlElement("w:t"); value.text="1"
    for node in (begin,inst,sep,value,end): run._r.append(node)
    font(run,8.5,color=MUTED)

doc=Document(); section=doc.sections[0]; section.page_width,section.page_height=Inches(8.5),Inches(11); section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1); section.header_distance=section.footer_distance=Inches(.492)
normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11); normal._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"),"Calibri"); normal.font.color.rgb=RGBColor.from_string(INK); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
for name,size,color,before,after in (("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",12,"1F4D78",10,5)):
    s=doc.styles[name]; s.font.name="Calibri"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); s._element.rPr.rFonts.set(qn("w:hAnsi"),"Calibri"); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
lb=doc.styles["List Bullet"]; lb.font.name="Calibri"; lb.font.size=Pt(11); lb.paragraph_format.left_indent=Inches(.375); lb.paragraph_format.first_line_indent=Inches(-.188); lb.paragraph_format.space_after=Pt(4); lb.paragraph_format.line_spacing=1.25
header=section.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(header.add_run("GulfInfraHub  |  RFQ Feature Guide"),8.5,True,MUTED)
footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(footer.add_run("Client reference  •  August 2026  |  "),8.5,color=MUTED); page_number(footer)

para(doc,"CLIENT REFERENCE",10,True,AMBER,18,WD_ALIGN_PARAGRAPH.CENTER)
para(doc,"RFQ Feature Guide",28,True,NAVY,4,WD_ALIGN_PARAGRAPH.CENTER)
para(doc,"Requests for Quotation: buyer, supplier and administration workflows",15,color=BLUE,after=20,align=WD_ALIGN_PARAGRAPH.CENTER)
callout(doc,"Purpose","Provide a complete, implementation-based reference for discovering, creating, managing, quoting, evaluating and awarding RFQs in GulfInfraHub.")
heading(doc,"RFQ capability at a glance")
grid(doc,[("Buyer","Create an RFQ, save or publish it, manage its lifecycle, review quotations, chat with suppliers and award one offer."),("Supplier","Browse active RFQs, review requirements and documents, save or submit one quotation, track decisions, withdraw an active offer and communicate with the buyer."),("Administrator","Moderate RFQ publication status and restrict blocked accounts."),("Platform","Enforces ownership, closing dates, OTP authorization, quotation status transitions and an auditable event timeline.")])
heading(doc,"Guide structure",2)
para(doc,"1. Access and roles  •  2. RFQ marketplace  •  3. Creating an RFQ  •  4. Buyer workspace  •  5. Supplier quotations  •  6. Evaluation and award  •  7. Messaging and notifications  •  8. Security and administration  •  9. Status reference")

doc.add_page_break(); heading(doc,"1. Access points and user roles")
grid(doc,[("Main navigation","RFQs opens the marketplace at /rfqs."),("Homepage","Featured RFQ cards provide View RFQ and Submit Quote actions; homepage/global search can also locate published RFQs."),("Add Listing","Request for Quotation redirects to /rfqs?create=1 and opens the dedicated four-step creation modal."),("Buyer workspace","My RFQs lists RFQs owned by the signed-in buyer and their quotations."),("Supplier workspace","My Quotations lists offers created by the signed-in supplier."),("Messages / Notifications","Account destinations consolidate quotation conversations and event activity.")])
heading(doc,"Role permissions",2)
grid(doc,[("Visitor","May browse published, unexpired RFQs and view their public details; must sign in to submit a quotation or create an RFQ."),("Buyer / owner","May create and edit owned RFQs, publish or close them, review responses and change active quotation statuses."),("Supplier","May create one quotation per RFQ, edit it while Draft or Submitted, withdraw it while active, and access its conversation."),("Administrator","May approve, hold or reject eligible RFQs through moderation; cannot feature RFQs."),("Blocked user","Cannot create/edit RFQs, submit quotations, change statuses or send quotation messages.")])
callout(doc,"Separation of duties","The RFQ owner cannot quote on their own RFQ. Only the RFQ owner can make supplier-selection decisions.",True)

doc.add_page_break(); heading(doc,"2. RFQ marketplace and discovery")
para(doc,"The RFQ marketplace displays active procurement opportunities and allows users to refine the visible records immediately in the browser.")
heading(doc,"Search",2)
grid(doc,[("Search box","Case-insensitive partial search across RFQ title, project name, reference, city and country."),("Example keywords","ready-mix concrete; structural steel; equipment rental; RFQ-0007; Dubai; Riyadh; project name."),("Combined matching","The keyword and every selected filter must match the record."),("No result","A no-match message provides Clear all filters.")])
heading(doc,"Filters",2)
grid(doc,[("RFQ status","Multi-select checkboxes for Published, Draft and Closed, with record counts."),("Category","All Categories, Materials Sourcing, Equipment Rentals or Equipment Purchases."),("Country","All Countries or a country represented in the loaded RFQ records."),("City","All Cities or a city valid for the selected country; changing country resets city."),("Closing date","Exact match against the RFQ closing date."),("Reset","Clears statuses, category, country, city, closing date and keyword.")])
heading(doc,"RFQ card information and actions",2)
grid(doc,[("Identity","Status, requirement title, RFQ reference, project and category."),("Location / requirement","City, country, description, quantity, budget and closing date."),("Activity","Posted date, bid count and delivery terms."),("Details","Opens a modal with procurement, requirement, supporting-document and contact information."),("View / Quote","Opens the dedicated RFQ page and quotation area.")])
callout(doc,"Public listing rule","The standard marketplace query returns only Published RFQs whose expiration date is in the future.")

doc.add_page_break(); heading(doc,"3. Creating an RFQ")
para(doc,"Selecting Create New RFQ launches a four-stage form. The buyer may move backward before saving and receives field-level feedback if required data is missing or invalid.")
heading(doc,"Step 1 — Project & Procurement",2)
grid(doc,[("Project name","Required associated project."),("RFQ category","Required: Materials Sourcing, Equipment Rentals or Equipment Purchases."),("Material / Service","Required; becomes the RFQ title."),("GCC country","Required: Saudi Arabia, United Arab Emirates, Qatar, Kuwait, Oman or Bahrain."),("City","Required."),("Delivery date","Required and must be later than the closing date."),("Closing date","Required and used to stop new quotation acceptance."),("Estimated budget","Optional free-text commercial estimate.")],header=("Field","Requirement"))
heading(doc,"Step 2 — Requirements & Delivery",2)
grid(doc,[("Quantity","Required free-text quantity."),("Unit","Required; suggestions include Units, Tons, Kilograms, Liters, Meters, Square/Cubic Meters and Months."),("Specifications","Required; minimum 20 characters."),("Delivery address","Required; also stored as delivery terms."),("Notes","Optional commercial, inspection or other instructions."),("Corporate phone","Required."),("Procurement email","Required valid email; for a new RFQ it is locked to the signed-in account when available.")],header=("Field","Requirement"))

doc.add_page_break(); heading(doc,"3.1 Documents, review and authorization")
heading(doc,"Step 3 — Supporting Documents",2)
grid(doc,[("BOQ URL","Optional link to the bill of quantities."),("Drawings URL","Optional link to drawings."),("Specification-document URL","Optional link to a detailed technical specification."),("Other document URLs","Optional list separated by new lines or commas; each retained as a separate link.")])
callout(doc,"Document model","The current interface stores secure file URLs; it does not upload files directly from the RFQ form.",True)
heading(doc,"Step 4 — Review & Publish",2)
bullet(doc,"Review summary","Shows project/procurement data, requirements/delivery data, contact details and attached document links.")
bullet(doc,"Save Draft","Creates or updates the RFQ with Draft status.")
bullet(doc,"Publish RFQ","Creates or updates it with Published status and makes an eligible record discoverable.")
bullet(doc,"Reference generation","A new record receives an automatic sequential reference in RFQ-0001 format.")
heading(doc,"Email OTP authorization for a new RFQ",2)
grid(doc,[("Code","Four digits, sent to the procurement/account email."),("Rate limit","One OTP request per purpose/email within one minute."),("Validity","Expires after 10 minutes."),("Attempts","Five failed attempts invalidate the code flow and require a new code."),("Delivery","Email delivery is required for RFQ creation."),("Single use","A successfully verified RFQ-creation authorization is claimed once when the new RFQ is saved; reuse is rejected."),("Session","Successful verification signs the user in and marks the email verified; a new account is prompted to create a password.")])
callout(doc,"Editing exception","Editing an owned RFQ does not repeat the new-record OTP claim, but normal sign-in, ownership and blocked-account checks still apply.")

doc.add_page_break(); heading(doc,"3.2 Validation and editing rules")
grid(doc,[("Required core data","Project, requirement, category, GCC country, city, quantity, unit, address, dates, phone and valid email."),("Technical detail","Specifications must contain at least 20 characters."),("Date sequence","Delivery must be after the RFQ closing date."),("Ownership","An RFQ can be edited only by its buyer/owner."),("Terminal records","Awarded or Cancelled RFQs cannot be edited."),("Verification mismatch","A new RFQ is rejected if the submitted email differs from the verified account email."),("Expired/used authorization","The buyer is instructed to request a new OTP."),("Failure handling","Validation returns highlighted fields; unexpected persistence failure returns a retry message.")])
heading(doc,"Stored RFQ information",2)
para(doc,"In addition to visible fields, the platform stores the owner, unique reference, status, posted and updated timestamps, quotation relationship and awarded timestamp. Urgency is currently saved as “Medium Urgency.”")
heading(doc,"Recommended buyer preparation",2)
bullet(doc,"Define scope","Use measurable quantities, units, standards, grades and acceptance requirements.")
bullet(doc,"Align dates","Allow enough response time before the closing date and schedule delivery afterward.")
bullet(doc,"Provide usable documents","Confirm external links are authorized, accessible and current.")
bullet(doc,"State commercial expectations","Add budget only if appropriate and use Notes for inspections, logistics and payment context.")

doc.add_page_break(); heading(doc,"4. Buyer workspace — My RFQs")
grid(doc,[("Search","Matches owned RFQs by title, reference or project name."),("Status filter","All, Draft, Published, Closed, Awarded or Cancelled."),("Post new RFQ","Opens the creation workflow."),("RFQ summary","Shows reference, category, title, location, closing date and current status."),("View","Opens the RFQ detail page."),("Edit","Available unless the RFQ is Awarded or Cancelled."),("Publish","Available for Draft records."),("Close RFQ","Available for Published records."),("Quotation panel","Displays each supplier/company, total offer, delivery lead time, warranty and quotation status.")])
heading(doc,"Buyer lifecycle controls",2)
bullet(doc,"Draft → Published","The Publish control makes a valid draft available to marketplace discovery.")
bullet(doc,"Published → Closed","Close RFQ prevents the opportunity from accepting quotations and removes it from the normal active browse query.")
bullet(doc,"Any editable non-terminal record","The owner can update details and choose Draft or Published in the review step.")
bullet(doc,"Awarded","Set automatically after a supplier is awarded; records the award date.")
callout(doc,"Current UI scope","The buyer status control exposes Publish and Close. Cancelled is recognized by the data model/workspace but no buyer Cancel button is implemented in the reviewed interface.",True)

doc.add_page_break(); heading(doc,"5. Supplier quotation workflow")
heading(doc,"Eligibility",2)
grid(doc,[("Authentication","The supplier must be signed in and not blocked."),("RFQ state","The RFQ must be Published and its closing time must still be in the future."),("Ownership conflict","The RFQ buyer cannot submit a quotation to their own request."),("Quotation count","The database enforces one quotation per supplier per RFQ; a Draft/Submitted offer is edited rather than duplicated."),("Edit window","Only Draft and Submitted quotations are editable.")])
heading(doc,"Quotation fields",2)
grid(doc,[("Company / contact","Company name, contact person, valid email and phone — all required."),("Commercial","Unit price, total price and currency — all required. Currency options: AED, SAR, QAR, KWD, OMR, BHD and USD."),("Delivery","Required delivery time/lead time, such as 14 days."),("Warranty","Required."),("Offer valid until","Required for submission; optional while saving a draft."),("Payment terms","Required."),("Technical specification","Required; minimum 20 characters."),("Remarks","Optional supplier notes."),("Quotation PDF URL","Optional link to a quotation document.")],header=("Field group","Details"))
heading(doc,"Save and submit",2)
bullet(doc,"Save draft","Stores Draft status and adds a “Quotation saved as draft” event.")
bullet(doc,"Submit quotation","Stores Submitted status, submission timestamp and a buyer-visible event.")
bullet(doc,"Update submitted offer","A supplier may revise a Submitted offer while the RFQ remains open; the action creates another status event.")
bullet(doc,"Offer amount display","The platform combines currency and total price for list and buyer views.")

doc.add_page_break(); heading(doc,"5.1 My Quotations and detailed offer view")
grid(doc,[("Quotation list","Shows RFQ reference/title, buyer identity, total offer and quotation status."),("Browse RFQs","Returns to the marketplace."),("View","Opens the quotation details, timeline and conversation."),("Edit","Shown while the quotation is Draft or Submitted; opens the RFQ quotation form."),("Withdraw","Shown for Submitted, Under Review or Shortlisted offers."),("Empty state","Provides a route to Find an RFQ to quote.")])
heading(doc,"Detailed quotation view",2)
grid(doc,[("Commercial summary","Unit price, total, delivery, warranty, payment terms and valid-until date."),("Technical offer","Displays the submitted specification and a Download quotation PDF action when a URL exists."),("Status badge","Shows Draft, Submitted, Under Review, Shortlisted, Awarded, Rejected or Withdrawn as applicable."),("Timeline","Lists recorded status, explanatory note and date/time."),("Conversation","Shows buyer/supplier messages and linked attachments."),("Award actions","For Awarded offers, provides an award-letter download and direct follow-up options.")])
heading(doc,"Withdrawal",2)
para(doc,"A supplier can withdraw only an active Submitted, Under Review or Shortlisted offer. The platform sets Withdrawn status, stores the withdrawal time, creates a timeline event and refreshes all affected workspaces. Withdrawal is not offered for Draft, Awarded, Rejected or already Withdrawn records.")

doc.add_page_break(); heading(doc,"6. Buyer evaluation and award")
para(doc,"The buyer can open Details & chat from My RFQs or the quotation detail page, compare the commercial/technical submission and move active offers through evaluation statuses.")
grid(doc,[("Under Review","Marks an active offer as being evaluated."),("Shortlisted","Identifies an offer retained for final comparison."),("Reject","Ends consideration for that quotation."),("Award supplier","Selects the quotation as the winner and completes the RFQ award workflow.")])
heading(doc,"Award transaction",2)
bullet(doc,"Winning quotation","Changes to Awarded and receives a buyer-generated timeline event.")
bullet(doc,"Competing quotations","All other active Submitted, Under Review or Shortlisted quotations for that RFQ are changed to Rejected.")
bullet(doc,"Competitor events","Each automatically rejected active quotation receives the note “Another supplier was awarded this RFQ.”")
bullet(doc,"RFQ record","Changes to Awarded and records the award timestamp.")
bullet(doc,"Atomic update","The winning decision, competing-offer rejections and RFQ update occur in one database transaction.")
callout(doc,"Decision boundary","Buyer status changes apply only to active quotations and only when the signed-in user owns the related RFQ.")
heading(doc,"Award letter",2)
para(doc,"The buyer and winning supplier can download a plain-text award letter. It identifies the date, RFQ reference, project, requirement, buyer, supplier, awarded amount, delivery time, warranty and payment terms. It records the marketplace decision but explicitly states that the parties must execute their final commercial agreement separately.")

doc.add_page_break(); heading(doc,"7. Messaging, notifications and audit trail")
heading(doc,"Buyer–supplier messaging",2)
grid(doc,[("Conversation scope","A separate conversation is attached to each quotation."),("Access","Only the quotation supplier and the related RFQ buyer can read or post in it."),("Message content","Text up to 3,000 characters."),("Attachment","Optional file URL; an attachment may be sent without text, in which case “Shared a file” is recorded."),("Messages page","Lists accessible conversations with RFQ reference/title, latest message preview and date."),("Detailed view","Shows sender, body, attachment link and timestamp.")])
heading(doc,"Notifications",2)
para(doc,"The Notifications page lists quotation events for RFQs where the user is either the supplier or buyer. Each entry shows the RFQ reference, event status, RFQ title, event note and date, and links to the quotation workspace.")
heading(doc,"Recorded events",2)
grid(doc,[("Draft / Submitted","Created whenever a supplier saves or submits an offer."),("Under Review / Shortlisted / Rejected / Awarded","Created when the buyer changes an active quotation status."),("Withdrawn","Created when the supplier withdraws an active quotation."),("Automatic rejection","Created for competing active offers when another supplier is awarded.")])
callout(doc,"Legacy records","If a quotation predates event tracking, its timeline may display that no events were recorded.")

doc.add_page_break(); heading(doc,"8. Authentication, authorization and security")
grid(doc,[("Signed-in session","Protected actions resolve the current user before operating."),("Email verification","A fresh RFQ-purpose OTP is required before every new RFQ save; the submitted email must match the verified account."),("Single-use authorization","The consumed OTP is claimed within the RFQ creation transaction and cannot authorize a second RFQ."),("Ownership","RFQ edits/status changes require buyer ownership; quotation edits/withdrawals require supplier ownership."),("Conversation privacy","Only the linked buyer and supplier can access quotation messages."),("Public visibility","Non-published RFQ detail pages are hidden from non-owners."),("Closing enforcement","Quotation saves are rejected after expiration or whenever the RFQ is not Published."),("Blocked accounts","Blocked users are denied RFQ, quotation, status and message mutations."),("Data relationships","Deleting an RFQ cascades to its quotations; deleting a quotation cascades to messages/events. Removed users are detached from owned RFQ/quotation records where configured.")])
heading(doc,"8.1 Administration and moderation",2)
grid(doc,[("Moderation scope","RFQs appear in the administrator listing-moderation workflow."),("Approve","Sets listing status to Published."),("Hold","Sets listing status to On Hold."),("Reject","Sets listing status to Rejected."),("Draft rule","Draft RFQs cannot be moderated."),("Feature control","RFQs are not eligible for the featured-listing action."),("Account control","An administrator can block a user, preventing further RFQ-related activity.")])
callout(doc,"Visibility note","The public RFQ marketplace expects Published status. Administrative Hold or Rejected records are therefore excluded from normal active discovery.")

doc.add_page_break(); heading(doc,"9. Status reference")
heading(doc,"RFQ statuses",2)
grid(doc,[("Draft","Owned work-in-progress; can be published."),("Published","Active marketplace record; accepts quotations until its expiration date."),("Closed","Buyer-closed record; no longer accepts quotations."),("Awarded","A supplier was selected; terminal for editing."),("Cancelled","Recognized terminal state; no buyer-facing Cancel control is currently exposed."),("On Hold / Rejected","Administrative moderation outcomes; not included in the My RFQs status selector or public active list.")],header=("Status","Meaning"))
heading(doc,"Quotation statuses",2)
grid(doc,[("Draft","Supplier work-in-progress; editable."),("Submitted","Sent to buyer; editable by supplier and eligible for buyer evaluation."),("Under Review","Buyer is evaluating; supplier may withdraw but cannot edit."),("Shortlisted","Buyer retained for final consideration; supplier may withdraw but cannot edit."),("Rejected","No longer under consideration; terminal in the UI."),("Awarded","Winning quotation; enables award letter and follow-up actions."),("Withdrawn","Supplier removed an active offer from consideration."),("Pending","Database default for legacy/direct records; the current form explicitly saves Draft or Submitted.")],header=("Status","Meaning"))
heading(doc,"End-to-end acceptance checklist",2)
for title,detail in (("Buyer creation","Create with OTP, verify reference and Draft/Published state."),("Marketplace discovery","Confirm published/unexpired visibility, search and all filters."),("Supplier response","Save Draft, submit, edit while Submitted and verify one-offer constraint."),("Buyer evaluation","Move through Under Review/Shortlisted and inspect events/notifications."),("Communication","Exchange text and attachment links from both authorized accounts."),("Award","Confirm winner, automatic competitor rejection, RFQ Awarded state and award letter."),("Security","Test non-owner, expired RFQ, blocked account and private conversation restrictions.")):
    bullet(doc,title,detail)

# Compact the two status-reference tables and final checklist so the reference
# closes on one balanced page without sacrificing content.
for table in doc.tables[-2:]:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs: run.font.size = Pt(9)
for paragraph in doc.paragraphs[-7:]:
    paragraph.paragraph_format.space_after = Pt(2); paragraph.paragraph_format.line_spacing = 1.05
    for run in paragraph.runs: run.font.size = Pt(9.5)

# Remove standalone page-break paragraphs and allow natural pagination. Dense
# reference tables then split cleanly without blank pages or orphaned headings.
paragraphs = list(doc.paragraphs)
for index, paragraph in enumerate(paragraphs[:-1]):
    if paragraph._p.xpath('.//w:br[@w:type="page"]'):
        paragraph._element.getparent().remove(paragraph._element)

OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
